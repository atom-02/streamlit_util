"""제목 + 링크 주소 + QR코드 안내문 생성기.

streamlit_app.py 에서 render() 를 호출해 사용합니다.
QR코드는 페이지 가로 중앙에 크게, 제목·링크는 왼쪽 정렬로 배치하며,
내용 블록이 몇 줄이든 위·아래 여백이 자동으로 균형 있게 맞춰진다.
"""
import io
import os
import re

import qrcode
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ===========================================================================
#  한글 폰트 등록 (Streamlit Community Cloud 배포 시 packages.txt 의
#  fonts-nanum 패키지로 아래 경로에 폰트가 설치된다)
# ===========================================================================
KOREAN_FONT_CANDIDATES = [
    ("KRFont", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf", None),
    ("KRFont-Bold", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf", None),
    ("KRFont", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 1),  # KR = index 1
    ("KRFont-Bold", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", 1),
    ("KRFont", "C:/Windows/Fonts/malgun.ttf", None),          # Windows에서 로컬 실행 시
    ("KRFont-Bold", "C:/Windows/Fonts/malgunbd.ttf", None),
    ("KRFont", "/System/Library/Fonts/Supplemental/AppleGothic.ttf", None),  # macOS
    ("KRFont-Bold", "/System/Library/Fonts/Supplemental/AppleGothic.ttf", None),
]


def register_korean_fonts():
    """사용 가능한 한글 폰트를 찾아 reportlab에 등록한다."""
    registered = {"regular": None, "bold": None}
    for name, path, index in KOREAN_FONT_CANDIDATES:
        slot = "regular" if name == "KRFont" else "bold"
        if registered[slot] is not None:
            continue
        if name not in pdfmetrics.getRegisteredFontNames():
            if not os.path.exists(path):
                continue
            try:
                if index is None:
                    pdfmetrics.registerFont(TTFont(name, path))
                else:
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=index))
            except Exception:
                continue
        registered[slot] = name

    if registered["regular"] is None:
        registered["regular"] = "Helvetica"
    if registered["bold"] is None:
        registered["bold"] = "Helvetica-Bold"
    return registered


FONTS = register_korean_fonts()


# ===========================================================================
#  공용 로직
# ===========================================================================
def normalize_url(url: str) -> str:
    url = url.strip()
    if url and not re.match(r"^[a-zA-Z][a-zA-Z0-9+\-.]*://", url):
        url = "https://" + url
    return url


def make_qr_image(url: str, box_size: int = 10, border: int = 2) -> Image.Image:
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=box_size,
        border=border,
    )
    qr.add_data(normalize_url(url))
    qr.make(fit=True)
    return qr.make_image(fill_color="black", back_color="white").convert("RGB")


def safe_filename(title: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", (title or "").strip() or "qr_code")


def _wrap_text_by_width(c, text, font_name, font_size, max_width):
    """긴 텍스트를 지정된 너비에 맞춰 여러 줄로 나눈다."""
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        if c.stringWidth(trial, font_name, font_size) > max_width and current:
            lines.append(current)
            current = ch
        else:
            current = trial
    if current:
        lines.append(current)
    return lines


# ===========================================================================
#  PDF 생성
# ===========================================================================
def create_qr_pdf_bytes(title: str, url: str) -> bytes:
    """제목 / 링크 주소 / QR코드가 담긴 1페이지 PDF를 만들어 bytes로 반환한다."""
    title = (title or "").strip() or "QR 코드"
    url = (url or "").strip()
    if not url:
        raise ValueError("링크 주소가 비어 있습니다.")
    url = normalize_url(url)

    page_w, page_h = A4
    left_margin = 37 * mm
    min_margin = 20 * mm

    title_font_size = 32
    title_line_height = title_font_size * 1.15
    label_font_size = 18
    url_font_size = 18
    gap_title_to_label = 14 * mm
    gap_label_to_url = 10 * mm
    url_line_height = 10 * mm
    gap_url_to_qr = 6 * mm

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    max_text_width = page_w - 2 * left_margin

    c.setFont(FONTS["bold"], title_font_size)
    title_lines = _wrap_text_by_width(c, title, FONTS["bold"], title_font_size, max_text_width)
    c.setFont(FONTS["regular"], url_font_size)
    url_lines = _wrap_text_by_width(c, url, FONTS["regular"], url_font_size, max_text_width)

    text_height = (
        len(title_lines) * title_line_height
        + gap_title_to_label
        + gap_label_to_url
        + len(url_lines) * url_line_height
    )

    desired_qr_size = 150 * mm
    available_for_qr = page_h - 2 * min_margin - text_height - gap_url_to_qr
    qr_size = min(desired_qr_size, page_w - 20 * mm, available_for_qr)
    qr_size = max(qr_size, 40 * mm)

    total_block_height = text_height + gap_url_to_qr + qr_size
    top_margin = max(min_margin, (page_h - total_block_height) / 2)

    y = page_h - top_margin

    c.setFont(FONTS["bold"], title_font_size)
    for line in title_lines:
        c.drawString(left_margin, y, line)
        y -= title_line_height

    y -= gap_title_to_label
    c.setFont(FONTS["bold"], label_font_size)
    c.drawString(left_margin, y, "링크 주소")

    y -= gap_label_to_url
    c.setFont(FONTS["regular"], url_font_size)
    for line in url_lines:
        c.drawString(left_margin, y, line)
        y -= url_line_height

    y -= gap_url_to_qr
    qr_img = make_qr_image(url)
    qr_x = (page_w - qr_size) / 2
    img_buf = io.BytesIO()
    qr_img.save(img_buf, format="PNG")
    img_buf.seek(0)
    from reportlab.lib.utils import ImageReader
    y -= qr_size
    c.drawImage(ImageReader(img_buf), qr_x, y, width=qr_size, height=qr_size)
    c.save()

    return buf.getvalue()


# ===========================================================================
#  워드(.docx) 생성
# ===========================================================================
def _set_table_borderless(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _add_hyperlink(paragraph, url, text, color="0563C1", underline=True):
    """단락에 클릭 가능한 하이퍼링크를 추가한다.

    color/underline 을 지정하지 않으면 기본 하이퍼링크 스타일(파란색 밑줄)을
    사용하고, color="000000", underline=False 처럼 넘기면 일반 텍스트와
    똑같은 모양이면서 클릭은 되는 링크를 만들 수 있다.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        color_el = OxmlElement("w:color")
        color_el.set(qn("w:val"), color)
        rPr.append(color_el)
    if underline:
        underline_el = OxmlElement("w:u")
        underline_el.set(qn("w:val"), "single")
        rPr.append(underline_el)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_qr_docx_bytes(title: str, url: str) -> bytes:
    """제목 / 링크 주소 / QR코드가 담긴 1페이지 워드 문서를 만들어 bytes로 반환한다.

    여백·글자 크기·줄바꿈·QR 크기 계산을 PDF 버전(create_qr_pdf_bytes)과
    완전히 동일한 수치로 맞춰, 두 출력물이 같은 모양으로 보이도록 한다.
    (reportlab 의 포인트 단위와 python-docx 의 Pt 단위가 동일하므로 계산값을
    그대로 재사용할 수 있다.) 보이지 않는 표(1행 1열)를 페이지 전체 높이에
    걸쳐 배치하고 세로 가운데 정렬을 적용해, 내용이 몇 줄이든 위·아래 여백이
    항상 균형 있게 나온다 — 여백을 상하 대칭으로 두면 PDF의 "전체 블록을
    페이지 중앙에 배치" 효과와 동일해진다.
    """
    title = (title or "").strip() or "QR 코드"
    url = (url or "").strip()
    if not url:
        raise ValueError("링크 주소가 비어 있습니다.")
    url = normalize_url(url)

    # ---- PDF 버전과 동일한 레이아웃 상수 / 계산 ----
    page_w, page_h = A4
    left_margin = 37 * mm
    min_margin = 20 * mm

    title_font_size = 32
    title_line_height = title_font_size * 1.15
    label_font_size = 18
    url_font_size = 18
    gap_title_to_label = 14 * mm
    gap_label_to_url = 10 * mm
    url_line_height = 10 * mm
    gap_url_to_qr = 6 * mm

    max_text_width = page_w - 2 * left_margin

    # 줄바꿈 위치를 PDF와 동일하게 계산하기 위한 측정 전용 캔버스(출력되지 않음)
    measure_buf = io.BytesIO()
    mc = canvas.Canvas(measure_buf, pagesize=A4)
    title_lines = _wrap_text_by_width(mc, title, FONTS["bold"], title_font_size, max_text_width)
    url_lines = _wrap_text_by_width(mc, url, FONTS["regular"], url_font_size, max_text_width)

    text_height = (
        len(title_lines) * title_line_height
        + gap_title_to_label
        + gap_label_to_url
        + len(url_lines) * url_line_height
    )

    desired_qr_size = 150 * mm
    available_for_qr = page_h - 2 * min_margin - text_height - gap_url_to_qr
    qr_size = min(desired_qr_size, page_w - 20 * mm, available_for_qr)
    qr_size = max(qr_size, 40 * mm)

    # ---- 문서 생성 ----
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(page_w)
    section.page_height = Pt(page_h)
    section.left_margin = Pt(left_margin)
    section.right_margin = Pt(left_margin)
    section.top_margin = Pt(min_margin)
    section.bottom_margin = Pt(min_margin)

    usable_h = page_h - 2 * min_margin
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    row = table.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    row.height = Pt(usable_h)
    cell = table.cell(0, 0)
    cell.width = Pt(max_text_width)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_table_borderless(table)

    def _line_paragraph(text_line, bold, size, space_after_pt=0):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after_pt)
        run = p.add_run(text_line)
        run.bold = bold
        run.font.size = Pt(size)
        return p

    # 제목 (표의 기본 첫 문단을 첫 줄로 재사용)
    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(0)
    run_title = p_title.add_run(title_lines[0])
    run_title.bold = True
    run_title.font.size = Pt(title_font_size)
    for line in title_lines[1:]:
        _line_paragraph(line, True, title_font_size)
    cell.paragraphs[len(title_lines) - 1].paragraph_format.space_after = Pt(gap_title_to_label)

    # "링크 주소" 라벨
    _line_paragraph("링크 주소", True, label_font_size, gap_label_to_url)

    # 링크 주소 (줄바꿈은 PDF와 동일 · 클릭 가능한 링크는 유지하되
    # 모양은 PDF처럼 검은색 · 밑줄 없음으로 맞춤)
    for i, line in enumerate(url_lines):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(gap_url_to_qr if i == len(url_lines) - 1 else 0)
        _add_hyperlink(p, url, line, color="000000", underline=False)
        for run in p.runs:
            run.font.size = Pt(url_font_size)

    # QR 코드 (가로 중앙, 크기는 PDF와 동일한 방식으로 계산)
    p_qr = cell.add_paragraph()
    p_qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_qr.paragraph_format.space_before = Pt(0)
    qr_img = make_qr_image(url)
    img_buf = io.BytesIO()
    qr_img.save(img_buf, format="PNG")
    img_buf.seek(0)
    p_qr.add_run().add_picture(img_buf, width=Pt(qr_size))

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()


# ===========================================================================
#  Streamlit UI  (streamlit_app.py 에서 render() 호출)
# ===========================================================================
def render():
    import streamlit as st

    st.subheader("🔗 QR코드 안내문 생성기")
    st.caption("제목과 링크 주소를 입력하면, QR코드가 담긴 안내문을 PDF 또는 워드 파일로 만들어 드립니다.")

    title = st.text_input(
        "제목", value="", placeholder="예: 미적분 패들렛 QR코드", key="qr_title",
    )
    url = st.text_input(
        "링크 주소 (URL)", value="", placeholder="예: https://padlet.com/...", key="qr_url",
    )

    if not url.strip():
        st.info("링크 주소를 입력하면 QR 코드 미리보기와 다운로드 버튼이 나타납니다.")
        return

    norm_url = normalize_url(url)

    col_preview, col_info = st.columns([1, 2])
    with col_preview:
        try:
            preview_img = make_qr_image(norm_url, box_size=6)
            st.image(preview_img, caption="QR 코드 미리보기", width=180)
        except Exception as e:
            st.error(f"QR 코드를 만드는 중 오류가 발생했습니다.\n{e}")
            return
    with col_info:
        st.write(f"**제목**  \n{title.strip() or 'QR 코드'}")
        st.write(f"**링크 주소**  \n{norm_url}")

    base_name = safe_filename(title)

    dl_pdf, dl_docx = st.columns(2)
    with dl_pdf:
        try:
            pdf_bytes = create_qr_pdf_bytes(title, url)
            st.download_button(
                "⬇ PDF 다운로드", pdf_bytes, file_name=f"{base_name}.pdf",
                mime="application/pdf", type="primary",
                use_container_width=True, key="qr_pdf_dl",
            )
        except Exception as e:
            st.error(f"PDF 생성 실패: {e}")
    with dl_docx:
        try:
            docx_bytes = create_qr_docx_bytes(title, url)
            st.download_button(
                "⬇ 워드 다운로드", docx_bytes, file_name=f"{base_name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary", use_container_width=True, key="qr_docx_dl",
            )
        except Exception as e:
            st.error(f"워드 문서 생성 실패: {e}")


if __name__ == "__main__":
    print("이 파일은 모듈입니다. 다음처럼 실행하세요:\n"
          "    pip install -r requirements.txt\n"
          "    streamlit run streamlit_app.py")
